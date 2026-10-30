"""F-8b: the UI grader script is blueprint-derived and a skipped run fails."""

from __future__ import annotations

import json

import pytest

from app.factory.build.product_grader import ui_grader as ug

BPS = [
    {"product_id": "lettings-desk", "product_name": "Lettings Desk", "summary": "tenancy files",
     "ui_modules": ["tenant_portal"], "capabilities": [{"id": "rent_roll", "description": "Roll rent."}, {"id": "notice_log", "description": "Log notices."}]},
    {"product_id": "site-ops", "product_name": "Site Ops", "summary": "inspections",
     "ui_modules": ["command_center", "map"], "capabilities": [{"id": "defect_register", "description": "Raise defects."}]},
]


def _report(names, failed=()):
    return {"suites": [{"specs": [{"title": n, "tests": [{"results": [{"status": "failed" if n in failed else "passed"}]}]} for n in names]}]}


@pytest.mark.parametrize("bp", BPS)
def test_plan_covers_every_declared_route_and_screen(bp):
    names = [s.name for s in ug.ui_plan(bp)]
    for c in bp["capabilities"]:
        assert f"route:/v1/{c['id']}" in names
    for m in bp["ui_modules"]:
        assert f"screen:{m}" in names
    assert {s.kind for s in ug.ui_plan(bp)} >= set(ug.STEP_KINDS)


def test_two_blueprints_two_scripts():
    a, b = (ug.render_spec(bp) for bp in BPS)
    assert a != b
    assert "rent_roll" in a and "rent_roll" not in b


@pytest.mark.parametrize("bp", BPS)
def test_all_passed_is_ok(bp):
    g = ug.parse_report(_report([s.name for s in ug.ui_plan(bp)]), bp)
    assert g.ok


@pytest.mark.parametrize("bp", BPS)
def test_failed_or_missing_step_is_named(bp):
    names = [s.name for s in ug.ui_plan(bp)]
    g = ug.parse_report(_report(names[1:], failed=[names[-1]]), bp)
    assert not g.ok and g.failing == [names[0], names[-1]]


@pytest.mark.parametrize("base_url", [None, ""])
def test_grader_mutant_skip_ui_run_fails(tmp_path, base_url):
    g = ug.grade_ui(BPS[0], tmp_path, base_url=base_url)
    assert not g.ok and not g.ran


def test_run_without_report_fails(tmp_path):
    g = ug.grade_ui(BPS[1], tmp_path, base_url="http://x", run=lambda *a, **k: None)
    assert not g.ok


def test_run_with_report_is_parsed(tmp_path):
    bp = BPS[1]

    def fake(cmd, **kw):
        names = [s.name for s in ug.ui_plan(bp)]
        open(kw["env"]["PLAYWRIGHT_JSON_OUTPUT_NAME"], "w").write(json.dumps(_report(names)))

    g = ug.grade_ui(bp, tmp_path, base_url="http://x", run=fake)
    assert g.ok
    assert (tmp_path / "fixtures" / "fixture.pdf").read_bytes().startswith(b"%PDF")
    assert (tmp_path / "fixtures" / "fixture.xlsx").is_file()


@pytest.mark.parametrize("bp", BPS)
def test_fixtures_come_from_blueprint(tmp_path, bp):
    fx = ug.write_fixtures(bp, tmp_path)
    assert bp["product_name"].encode() in fx["pdf"].read_bytes()


# --- G2 survivors M1-M7 ---------------------------------------------------

@pytest.mark.parametrize("statuses,expected", [
    (["passed", "passed"], True),
    (["passed", "failed"], False),   # M1: any() would pass this
    (["failed", "passed"], False),
    ([], False),                     # M2: empty results never pass
])
def test_step_passed_needs_every_result(statuses, expected):
    assert ug.step_passed([{"status": x} for x in statuses]) is expected


@pytest.mark.parametrize("bp", BPS)
def test_report_with_mixed_retries_or_empty_step_fails(bp):
    names = [s.name for s in ug.ui_plan(bp)]
    rep = _report(names)
    specs = rep["suites"][0]["specs"]
    specs[0]["tests"][0]["results"].append({"status": "failed"})
    specs[1]["tests"][0]["results"] = []
    g = ug.parse_report(rep, bp)
    assert g.failing == [names[0], names[1]]


@pytest.mark.parametrize("status,honest", [(422, True), (503, True), (200, False), (500, False)])
def test_honest_error_never_accepts_200(status, honest):  # M3
    assert (status in ug.HONEST_ERROR_STATUSES) is honest
    assert f"expect({list(ug.HONEST_ERROR_STATUSES)}".replace(" ", "") in ug.KIND_JS["honest_error"].replace(" ", "")


@pytest.mark.parametrize("exc", [OSError("boom"), TimeoutError("slow"), RuntimeError("x")])
def test_run_exception_is_a_recorded_fail(tmp_path, exc):  # M4
    def boom(*a, **k):
        raise exc

    g = ug.grade_ui(BPS[0], tmp_path, base_url="http://x", run=boom)
    assert not g.ok and any(type(exc).__name__ in e for e in g.errors)


@pytest.mark.parametrize("kind,needle", [
    ("route", "toBe(200)"),                 # M5
    ("load", "toBe(200)"),
    ("screen", "toBe(200)"),
    ("screen", "expect(errs).toEqual([])"),  # M6
    ("load", "expect(errs).toEqual([])"),
])
def test_kind_assertions(kind, needle):
    assert needle in ug.KIND_JS[kind]


@pytest.mark.parametrize("bp", BPS)
def test_spec_has_a_body_for_every_planned_kind(bp):
    spec = ug.render_spec(bp)
    for kind in {s.kind for s in ug.ui_plan(bp)}:
        assert f'(s.kind === "{kind}")' in spec
        assert ug.KIND_JS[kind].splitlines()[0] in spec


@pytest.mark.parametrize("bp", BPS)
def test_xlsx_fixture_rows_come_from_blueprint(tmp_path, bp):  # M7
    from openpyxl import load_workbook

    ws = load_workbook(ug.write_fixtures(bp, tmp_path)["xlsx"]).active
    rows = [r for r in ws.iter_rows(values_only=True)][1:]
    assert [r[0] for r in rows] == [c["id"] for c in bp["capabilities"]]
    assert [r[1] for r in rows] == [c["description"] for c in bp["capabilities"]]
    assert ws.title == ug._slug(bp["product_id"])


def test_two_blueprints_two_xlsx(tmp_path):
    from openpyxl import load_workbook

    a, b = (load_workbook(ug.write_fixtures(bp, tmp_path / str(i))["xlsx"]).active for i, bp in enumerate(BPS))
    assert list(a.iter_rows(values_only=True)) != list(b.iter_rows(values_only=True))


# --- G2 U14-U17, U21 ------------------------------------------------------

@pytest.mark.parametrize("bp", BPS)
def test_spec_resolves_every_target_against_base_url(bp):  # U21
    spec = ug.render_spec(bp)
    assert "process.env.BASE_URL" in spec and "new URL(t, BASE)" in spec
    for call in ("page.goto(", "request.get(", "request.post("):
        for chunk in spec.split(call)[1:]:
            assert chunk.startswith("url("), call


@pytest.mark.parametrize("url", ["http://127.0.0.1:8000", "https://p.example.test/app/"])
def test_grade_ui_passes_base_url_to_the_run(tmp_path, url):
    seen = {}
    ug.grade_ui(BPS[0], tmp_path, base_url=url, run=lambda cmd, **kw: seen.update(kw["env"]))
    assert seen["BASE_URL"] == url


@pytest.mark.parametrize("kind,needles", [
    ("bind", ["s.target", "toBeVisible()"]),                                   # U14
    ("attach_pdf", ["name: '+'", "setInputFiles", "GRADE_FIXTURES"]),          # U15
    ("attach_xlsx", ["name: '+'", "setInputFiles", "GRADE_FIXTURES"]),
    ("parsed", ["/parsed/i", "toBeVisible"]),                                   # U16
    ("cited_answer", ["QUESTION", "citation", "getByText(/", "toBeVisible"]),           # U17
])
def test_kind_bodies_pinned(kind, needles):
    for n in needles:
        assert n in ug.KIND_JS[kind], (kind, n)


@pytest.mark.parametrize("bp", BPS)
def test_attach_targets_are_the_written_fixtures(tmp_path, bp):
    fx = ug.write_fixtures(bp, tmp_path)
    targets = {s.target for s in ug.ui_plan(bp) if s.kind.startswith("attach_")}
    assert targets == {fx["pdf"].name, fx["xlsx"].name}


# --- HTTP status asserts, executed for real (G2 U5, U12, U13, U15b, U16) ----
import sys as _sys  # noqa: E402
from pathlib import Path as _P  # noqa: E402

_sys.path.insert(0, str(_P(__file__).parent))
import ui_spec_harness as h  # noqa: E402

BP = BPS[0]
ROUTES = [f"/v1/{c['id']}" for c in BP["capabilities"]]
HONEST = f"POST /v1/{BP['capabilities'][0]['id']}"


@pytest.mark.parametrize("statuses,want", [
    ({}, {"load:/": "passed", "screen:tenant_portal": "passed"}),
    ({"/": 500}, {"load:/": "failed", "screen:tenant_portal": "failed"}),
    ({"/": 201}, {"load:/": "failed", "screen:tenant_portal": "failed"}),
    ({"/other": 500}, {"load:/": "passed", "screen:tenant_portal": "passed"}),
])
def test_page_steps_need_exactly_200(tmp_path, statuses, want):  # U5 / U12
    # A URL fragment never reaches the server, so the screen goes to "/".
    with h.server(statuses) as base:
        got = h.run_spec(ug.render_spec(BP), tmp_path, base, only_kinds=("load", "screen"))
    assert {k: got[k].split(":")[0] for k in want} == want
    if want.get("load:/") == "failed":
        assert f"expected 200 got {statuses['/']}" in got["load:/"]


@pytest.mark.parametrize("bad,code", [(ROUTES[0], 404), (ROUTES[1], 500), (ROUTES[1], 204)])
def test_every_route_needs_exactly_200(tmp_path, bad, code):  # U13
    with h.server({bad: code}) as base:
        got = h.run_spec(ug.render_spec(BP), tmp_path, base, only_kinds=("route",))
    assert got == {f"route:{r}": ("passed" if r != bad else f"failed: expected 200 got {code}") for r in ROUTES}


@pytest.mark.parametrize("code,ok", [(422, True), (503, True), (200, False), (400, False), (500, False)])
def test_honest_error_needs_422_or_503_never_200(tmp_path, code, ok):  # U15b
    with h.server({HONEST: code}) as base:
        got = h.run_spec(ug.render_spec(BP), tmp_path, base, only_kinds=("honest_error",))
    (v,) = got.values()
    assert (v == "passed") is ok


@pytest.mark.parametrize("pair", [({}, {ROUTES[0]: 500}), ({ROUTES[1]: 404}, {})])
def test_two_base_urls_two_runs(tmp_path, pair):  # U16: BASE_URL is what is graded
    runs = []
    for i, statuses in enumerate(pair):
        (tmp_path / str(i)).mkdir()
        with h.server(statuses) as base:
            runs.append(h.run_spec(ug.render_spec(BP), tmp_path / str(i), base, only_kinds=("route",)))
    assert runs[0] != runs[1]
    for statuses, run in zip(pair, runs):
        assert sorted(k for k, v in run.items() if v != "passed") == sorted(f"route:{r}" for r in statuses)


@pytest.mark.parametrize("base", [None, ""])
def test_no_base_url_throws_not_relative(tmp_path, base):  # U16
    got = h.run_spec(ug.render_spec(BP), tmp_path, base, only_kinds=("route",))
    assert "BASE_URL is not set" in got["__error__"]


# --- G2 15bd9c1 gaps on split A: U5 nested suites, U12 bind, U13 target ---
@pytest.mark.parametrize("depth", [1, 3])
def test_nested_report_suites_are_read(depth):  # U5
    names = [s.name for s in ug.ui_plan(BP)]
    suite = {"specs": [{"title": n, "tests": [{"results": [{"status": "passed"}]}]} for n in names]}
    for _ in range(depth):
        suite = {"suites": [suite]}
    g = ug.parse_report({"suites": [suite]}, BP)
    assert g.ok and [s["passed"] for s in g.steps] == [True] * len(names)


@pytest.mark.parametrize("i,served,ok", [(0, 0, True), (1, 1, True), (0, 1, False), (1, 0, False)])
def test_bind_needs_the_products_own_id_on_the_page(tmp_path, i, served, ok):  # U12
    bp = BPS[i]
    with h.server({}, bodies={"/": f"<h1>tenant {BPS[served]['product_id']}</h1>"}) as base:
        got = h.run_spec(ug.render_spec(bp), tmp_path, base, only_kinds=("load", "bind"))
    assert got["load:/"] == "passed" and (got["bind:tenant"] == "passed") is ok


@pytest.mark.parametrize("i", [0, 1])
def test_honest_error_targets_the_products_first_route(tmp_path, i):  # U13
    bp = BPS[i]
    first = bp["capabilities"][0]["id"]
    with h.server({f"POST /v1/{first}": 422}) as base:  # every other POST answers 200
        got = h.run_spec(ug.render_spec(bp), tmp_path, base, only_kinds=("honest_error",))
        assert [k for k, _ in base.seen] == [f"/v1/{first}"]
    assert got == {f"honest_error:/v1/{first}": "passed"}


@pytest.mark.parametrize("which", [None, ""])
def test_missing_node_fails_never_skips(tmp_path, monkeypatch, which):
    monkeypatch.setattr(h.shutil, "which", lambda name: which or None)
    with pytest.raises(pytest.fail.Exception, match="never a skip"):
        h.run_spec(ug.render_spec(BP), tmp_path, "http://127.0.0.1:9")


# --- merged from split B (#83): N1, N2, N4-N6, U18, U19 ---
PLANNED = sorted({s.kind for bp in BPS for s in ug.ui_plan(bp)})


@pytest.mark.parametrize("drop", [("route",), ("screen",), ("honest_error",), ("load", "parsed"), ("attach_pdf", "bind", "cited_answer")])
def test_missing_kind_is_a_hard_error_naming_every_kind(monkeypatch, drop):  # N1 / N2
    monkeypatch.setattr(ug, "KIND_JS", {k: v for k, v in ug.KIND_JS.items() if k not in drop})
    with pytest.raises(ValueError) as e:
        ug.render_spec(BPS[1])
    assert str(e.value) == f"no UI assertion for step kinds: {sorted(drop)}"


@pytest.mark.parametrize("extra", ["unused_kind", "another"])
def test_extra_kind_bodies_are_not_an_error(monkeypatch, extra):  # N4
    monkeypatch.setattr(ug, "KIND_JS", {**ug.KIND_JS, extra: "expect(1).toBe(1);"})
    assert f"s.kind === {json.dumps(extra)}" in ug.render_spec(BPS[0])


@pytest.mark.parametrize("bp", BPS)
def test_every_planned_kind_has_exactly_one_branch(bp):  # N5
    spec = ug.render_spec(bp)
    for k in {s.kind for s in ug.ui_plan(bp)}:
        assert spec.count(f"(s.kind === {json.dumps(k)})") == 1
    assert spec.count("} else { throw new Error('no assertion for ' + s.kind); }") == 1


@pytest.mark.parametrize("bp", BPS)
def test_planned_kinds_all_have_bodies(bp):  # N6
    assert {s.kind for s in ug.ui_plan(bp)} <= set(ug.KIND_JS)


def _with_plan_kind(spec: str, name: str, kind: str) -> str:
    plan_line = next(line for line in spec.splitlines() if line.startswith("const PLAN = "))
    plan = json.loads(plan_line[len("const PLAN = "):-1])
    plan.append({"name": name, "kind": kind, "target": "/"})
    return spec.replace(plan_line, f"const PLAN = {json.dumps(plan)};")


@pytest.mark.parametrize("kind", ["mystery", "routee"])
def test_runtime_unknown_kind_throws_never_passes(tmp_path, kind):  # U18
    spec = _with_plan_kind(ug.render_spec(BPS[0]), f"odd:{kind}", kind)
    with h.server({}) as base:
        got = h.run_spec(spec, tmp_path, base, only_kinds=(kind, "route"))
    assert got[f"odd:{kind}"] == f"failed: no assertion for {kind}"
    assert all(v == "passed" for k, v in got.items() if k.startswith("route:"))


@pytest.mark.parametrize("kind,statuses,failed", [
    ("route", {"/v1/rent_roll": 500}, {"route:/v1/rent_roll"}),
    ("load", {"/": 500}, {"load:/"}),
    ("honest_error", {"POST /v1/rent_roll": 200}, {"honest_error:/v1/rent_roll"}),
])
def test_each_kind_runs_its_own_branch(tmp_path, kind, statuses, failed):  # U19
    with h.server(statuses) as base:
        got = h.run_spec(ug.render_spec(BPS[0]), tmp_path, base, only_kinds=(kind,))
    assert {k for k, v in got.items() if v != "passed"} == failed and got


# --- G2 15bd9c1 gaps on split B: N4-N6, U18, U19 executed end to end -------
def _flow(tmp_path, bp, *, page="", upload="parsed ok", answer="total 1,240 [1]"):
    fx = ug.write_fixtures(bp, tmp_path / "fx")
    with h.server({}, bodies={"/": page, "POST /__upload": upload, "POST /__ask": answer}) as base:
        got = h.run_spec(ug.render_spec(bp), tmp_path, base, fixtures=fx["pdf"].parent,
                         only_kinds=("load", "bind", "attach_pdf", "attach_xlsx", "parsed", "cited_answer"))
        return got, list(base.seen), fx


@pytest.mark.parametrize("bp", BPS)
def test_full_flow_passes_on_a_product_that_answers(tmp_path, bp):
    got, seen, fx = _flow(tmp_path, bp, page=f"tenant {bp['product_id']}")
    assert set(got.values()) == {"passed"}, got
    uploads = [(k, b) for k, b in seen if k.startswith("/__upload")]
    assert [k for k, _ in uploads] == ["/__upload?name=fixture.pdf", "/__upload?name=fixture.xlsx"]  # N6
    assert uploads[0][1] == fx["pdf"].read_bytes() and uploads[1][1] == fx["xlsx"].read_bytes()
    assert bp["summary"].encode() in uploads[0][1]  # U19: the blueprint's own summary reaches the product
    asked = [b.decode() for k, b in seen if k == "/__ask"]
    assert asked == [ug._question(bp)]  # U18
    assert bp["capabilities"][0]["description"].split(".")[0] in asked[0]


@pytest.mark.parametrize("bp", BPS)
def test_bind_fails_without_the_product_id(tmp_path, bp):  # N4
    got, _, _ = _flow(tmp_path, bp, page="tenant somebody-else")
    assert got["bind:tenant"].startswith("failed")


@pytest.mark.parametrize("upload", ["", "upload received"])
def test_parsed_fails_without_a_parsed_status(tmp_path, upload):  # N5
    got, _, _ = _flow(tmp_path, BPS[0], page="lettings-desk", upload=upload)
    assert got["parsed:status"].startswith("failed") and got["attach:pdf"] == "passed"


@pytest.mark.parametrize("answer,ok", [("total 1,240 [1]", True), ("I cannot say", False), ("see [2]", True)])
def test_cited_answer_needs_a_citation_and_a_number(tmp_path, answer, ok):
    got, _, _ = _flow(tmp_path, BPS[1], page="site-ops", answer=answer)
    assert (got["answer:cited"] == "passed") is ok


@pytest.mark.parametrize("pair", [(0, 1)])
def test_two_blueprints_two_questions_and_fixtures(tmp_path, pair):  # U18 / U19 two inputs
    runs = [_flow(tmp_path / str(i), BPS[i], page=BPS[i]["product_id"]) for i in pair]
    asked = [[b for k, b in seen if k == "/__ask"] for _, seen, _ in runs]
    pdfs = [[b for k, b in seen if k == "/__upload?name=fixture.pdf"] for _, seen, _ in runs]
    assert asked[0] != asked[1] and pdfs[0] != pdfs[1]



# --- 'playwright report has no results' is an error, and only for an empty report ---
@pytest.mark.parametrize("report,empty", [
    ({"suites": []}, True),
    ({"suites": [{"specs": []}]}, True),
    ({}, True),
    ({"suites": [{"specs": [{"title": "load:/", "tests": [{"results": [{"status": "passed"}]}]}]}]}, False),
])
def test_empty_report_is_a_named_error(report, empty):
    g = ug.parse_report(report, BPS[0])
    assert g.errors == (["playwright report has no results"] if empty else [])
    assert ("playwright report has no results" in g.to_dict()["failing"]) is empty and not (empty and g.ok)
