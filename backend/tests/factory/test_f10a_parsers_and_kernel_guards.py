"""F-10a: F-1b parser pins + F-1c kernel guards (K-F1, done here).

* CSV: row numbers are the file's own (1-based, header = 1); a blank line and
  an all-blank-cells row (",,") are skipped without renumbering later rows.
* DOCX: a heading is decided by style IDENTITY (built-in + style_id), never by
  the display name: a custom style NAMED "Heading" is not a heading; a
  built-in heading whose display name is localized still is.
* Kernel path: every citation carries the retrieved chunk's file + location.
* K-F1: a definition's guards are evaluated with safe_eval BEFORE the
  expression; a failed guard is an UNKNOWN refusal naming the guard, and the
  answer route still returns 200 (never 500). Two inputs, two results.
"""

from __future__ import annotations

import io
from types import SimpleNamespace
from typing import List

import pytest

from app.cerebrum_product_kernel.answer import _kernel_answer, kernel_coverage
from app.cerebrum_product_kernel.formulas import answer_definition, load_base_definitions, resolve_definitions
from app.cerebrum_product_kernel.ingestion.parsers import parse_csv, parse_docx


# ── CSV ────────────────────────────────────────────────────────────────────────

def _csv_rows(data: bytes, name: str = "ledger.csv"):
    chunks, meta = parse_csv(data, source_name=name, object_id="obj", tenant_id="t")
    return [(c["location"]["row"], c["text"]) for c in chunks], meta


def test_csv_row_numbers_are_the_files_own():
    rows, meta = _csv_rows(b"item,qty\nbolt,4\nnut,7\n")
    assert [r for r, _ in rows] == [1, 2, 3]
    assert meta == {"rows": 3}
    rows_b, _ = _csv_rows(b"item,qty\nnut,7\n")
    assert [r for r, _ in rows_b] == [1, 2]
    assert "nut" in dict(rows_b)[2] and "bolt" in dict(rows)[2]


@pytest.mark.parametrize(
    "data, expected_rows",
    [
        (b"item,qty\n\nbolt,4\n", [1, 3]),          # empty line
        (b"item,qty\n,,\nbolt,4\n", [1, 3]),        # all-blank cells
        (b"item,qty\n , \nbolt,4\n,\nnut,7\n", [1, 3, 5]),
    ],
)
def test_csv_blank_rows_skipped_without_renumbering(data, expected_rows):
    rows, meta = _csv_rows(data)
    assert [r for r, _ in rows] == expected_rows
    assert meta["rows"] == len(expected_rows)
    assert all(text.replace("|", "").strip() for _, text in rows)


def test_csv_location_names_the_file():
    rows_a, _ = _csv_rows(b"a,b\n", name="one.csv")
    chunks, _ = parse_csv(b"a,b\n", source_name="two.csv", object_id="o", tenant_id="t")
    assert chunks[0]["location"]["sheet"] == "two.csv"
    assert chunks[0]["file"] == "two.csv"


# ── DOCX heading by style identity ──────────────────────────────────────────

def _docx_bytes(build) -> bytes:
    import docx

    document = docx.Document()
    build(document)
    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()


def _body_heading(data: bytes, body_text: str):
    chunks, _ = parse_docx(data, source_name="d.docx", object_id="o", tenant_id="t")
    body = [c for c in chunks if body_text in c["text"]]
    assert body, chunks
    return body[0]["location"]["heading"]


@pytest.mark.parametrize(
    "builtin, style_id, expected",
    [(True, "Heading1", True), (False, "Heading1", False), (True, "Heading", False)],
)
def test_heading_needs_builtin_and_a_heading_style_id(builtin, style_id, expected):
    """A real .docx cannot hold a custom style with a built-in heading id (ids
    are unique per style part), so the builtin flag is pinned on the predicate."""
    from app.cerebrum_product_kernel.ingestion.parsers import _is_heading_paragraph

    para = SimpleNamespace(style=SimpleNamespace(builtin=builtin, style_id=style_id, name="Heading 1"))
    assert _is_heading_paragraph(para) is expected


def test_custom_style_named_heading_is_not_a_heading():
    custom_name = "Heading"
    from docx.enum.style import WD_STYLE_TYPE

    def build(document):
        document.add_heading("Real Section", level=1)
        custom = document.styles.add_style(custom_name, WD_STYLE_TYPE.PARAGRAPH)
        document.add_paragraph("Fake Section", style=custom)
        document.add_paragraph("body after the custom style")

    assert _body_heading(_docx_bytes(build), "body after the custom style") == "Real Section"


def test_builtin_heading_with_localized_display_name_is_a_heading():
    def build(document):
        document.add_heading("Abschnitt Eins", level=1)
        document.styles["Heading 1"].name = "Überschrift 1"
        document.add_paragraph("body under the localized heading")

    data = _docx_bytes(build)
    import docx

    reloaded = docx.Document(io.BytesIO(data))
    assert reloaded.paragraphs[0].style.name != "Heading 1"
    assert _body_heading(data, "body under the localized heading") == "Abschnitt Eins"


# ── kernel path: citations + guards ─────────────────────────────────────────

class _Engine:
    """Returns fixed hits; the kernel path reads text, file and location."""

    def __init__(self, hits: List[SimpleNamespace]):
        self.hits = hits

    def retrieve(self, collection, query):
        return list(self.hits)


def _hit(text, file, location):
    return SimpleNamespace(text=text, file=file, location=location)


def test_kernel_citation_carries_file_and_location():
    row = answer_definition("gross_margin_ratio")
    a = _kernel_answer(
        _Engine([_hit("revenue is 100. cogs is 60.", "fin.csv", {"sheet": "fin.csv", "row": 4})]),
        "c", "gross margin ratio", row,
    )
    b = _kernel_answer(
        _Engine([_hit("revenue is 100. cogs is 60.", "plan.docx", {"paragraph": 9, "heading": "Margins"})]),
        "c", "gross margin ratio", row,
    )
    assert [c.to_dict() for c in a.citations] == [{"file": "fin.csv", "location": {"sheet": "fin.csv", "row": 4}}]
    assert [c.to_dict() for c in b.citations] == [{"file": "plan.docx", "location": {"paragraph": 9, "heading": "Margins"}}]


GUARDED = [
    # (definition id, corpus that trips the guard, corpus that passes, value on pass)
    ("net_from_gross", "gross amount is 110. tax rate is -1.", "gross amount is 110. tax rate is 0.1.", 100.0),
    ("late_fee_simple",
     "overdue amount is 1000. annual rate is 0.12. days overdue is 30. days in year is 0.",
     "overdue amount is 1000. annual rate is 0.12. days overdue is 30. days in year is 360.", 10.0),
    ("break_even_revenue", "fixed costs is 5000. contribution margin ratio is 0.",
     "fixed costs is 5000. contribution margin ratio is 0.25.", 20000.0),
    ("contribution_margin_ratio", "unit price is 0. unit variable cost is 3.",
     "unit price is 8. unit variable cost is 6.", 0.25),
]


@pytest.mark.parametrize("ident, bad, good, value", GUARDED)
def test_guard_refuses_by_name_and_passes_when_it_holds(ident, bad, good, value):
    row = answer_definition(ident)
    guards = resolve_definitions()[ident]["guards"]
    assert row is not None and guards, f"{ident} must declare a guard"
    assert "guards" not in row, "the index row is not widened (Q2 ruling)"
    refused = _kernel_answer(_Engine([_hit(bad, "f.md", {"line": 1})]), "c", ident, row)
    assert refused.unknown is True
    assert refused.kernel_value is None
    for guard in guards:
        assert guard in refused.text
    assert "ZeroDivisionError" not in refused.text
    ok = _kernel_answer(_Engine([_hit(good, "f.md", {"line": 1})]), "c", ident, row)
    assert ok.unknown is False
    assert ok.kernel_value == pytest.approx(value)


def test_guard_read_from_resolved_set_present_vs_absent():
    """Same row, same zero-divisor corpus; only the resolved definition's
    guards differ. Present -> refusal naming the guard; absent -> the
    expression's own ZeroDivisionError refusal (still UNKNOWN, no guard named).
    A guard that is false without any division also refuses."""
    base = resolve_definitions()
    row = answer_definition("contribution_margin_ratio")
    eng = _Engine([_hit("unit price is 0. unit variable cost is 3.", "f", {})])
    present = _kernel_answer(eng, "c", "x", row, resolved=base)
    absent_set = dict(base)
    absent_set["contribution_margin_ratio"] = {**base["contribution_margin_ratio"], "guards": []}
    absent = _kernel_answer(eng, "c", "x", row, resolved=absent_set)
    assert present.unknown and "guard failed: unit_price != 0" in present.text
    assert absent.unknown and "guard failed" not in absent.text

    strict = dict(base)
    strict["contribution_margin_per_unit"] = {
        **base["contribution_margin_per_unit"], "guards": ["unit_price > unit_variable_cost"]}
    row2 = answer_definition("contribution_margin_per_unit")
    neg = _kernel_answer(_Engine([_hit("unit price is 5. unit variable cost is 9.", "f", {})]), "c", "x", row2, resolved=strict)
    pos = _kernel_answer(_Engine([_hit("unit price is 9. unit variable cost is 5.", "f", {})]), "c", "x", row2, resolved=strict)
    plain = _kernel_answer(_Engine([_hit("unit price is 5. unit variable cost is 9.", "f", {})]), "c", "x", row2, resolved=base)
    assert neg.unknown is True and "unit_price > unit_variable_cost" in neg.text
    assert pos.unknown is False and pos.kernel_value == pytest.approx(4.0)
    assert plain.unknown is False and plain.kernel_value == pytest.approx(-4.0)


def test_contribution_margin_ratio_is_defined_and_consumable():
    base = load_base_definitions()
    ids = {d["id"] for d in base["definitions"]}
    assert "contribution_margin_ratio" in ids
    consumer = next(d for d in base["definitions"] if d["id"] == "break_even_revenue")
    assert "contribution_margin_ratio" in consumer["inputs"]
    assert kernel_coverage("what is the contribution margin ratio")["id"] == "contribution_margin_ratio"


# ── route: zero divisor -> UNKNOWN, 200 ─────────────────────────────────────

class _FlatEmbedder:
    semantic = True

    def __call__(self, text: str) -> List[float]:
        return [1.0, 0.0]


@pytest.mark.parametrize(
    "ratio, status_unknown, value",
    [("0", True, None), ("0.25", False, 20000.0)],
)
def test_answer_route_zero_divisor_is_unknown_200(tmp_path, monkeypatch, ratio, status_unknown, value):
    from fastapi import FastAPI
    from fastapi.testclient import TestClient

    from app.cerebrum_product_kernel.ingestion import object_store, parse_job, retrieval_index
    from app.cerebrum_product_kernel.ingestion.parsers import parse_object
    from app.cerebrum_product_kernel.ingestion.router import router as ingestion_router
    from app.cerebrum_product_kernel.retrieval_engine import InMemoryVectorStore, RetrievalEngine
    from tests.factory.test_f1c_answer_routes import _stub_tenancy

    monkeypatch.setenv("STORAGE_PATH", str(tmp_path / "storage"))
    monkeypatch.setenv("FACTORY_OUTPUTS_ROOT", str(tmp_path / "factory_outputs"))
    engine = RetrievalEngine(embedder=_FlatEmbedder(), store=InMemoryVectorStore())
    retrieval_index.set_engine_provider(lambda identity: engine)
    try:
        _stub_tenancy(monkeypatch, {"tok": "tenant-g"})
        data = f"Break-even inputs.\n\nFixed costs is 5000. Contribution margin ratio is {ratio}.\n".encode()
        stored = object_store.save_object("tenant-g", "be.md", data)
        job = parse_job.create_job("tenant-g", object_id=stored.object_id, source_name=stored.source_name)
        assert parse_object("tenant-g", stored.object_id, job_id=job.job_id).status is parse_job.ParseStatus.PARSED
        app = FastAPI()
        app.include_router(ingestion_router)
        resp = TestClient(app).post(
            "/v1/product/ingestion/answer",
            headers={"Authorization": "Bearer tok"},
            json={"query": "what is the break even revenue", "answer_text": "It is 1."},
        )
    finally:
        retrieval_index.set_engine_provider(None)
    assert resp.status_code == 200, resp.text
    body = resp.json()
    assert body["kind"] == "kernel"
    assert body["unknown"] is status_unknown
    if value is None:
        assert "contribution_margin_ratio != 0" in body["text"]
        assert body.get("kernel_value") is None
    else:
        assert body["kernel_value"] == pytest.approx(value)


# ── G6 FAIL e215ecd repair 1: raising guards (M12), overlay guards on the
#    answer_query path (M14 resolved= dropped, M15 overlays= dropped) ────────

@pytest.mark.parametrize(
    "guard",
    [
        "unit_count != 0",            # names an input the definition never grounds -> MissingInput
        "unit_price / (unit_price - unit_price) > 0",  # divides by zero inside the guard
    ],
)
def test_guard_that_raises_is_an_unknown_refusal_naming_it(guard):
    base = resolve_definitions()
    strict = dict(base)
    strict["contribution_margin_per_unit"] = {**base["contribution_margin_per_unit"], "guards": [guard]}
    row = answer_definition("contribution_margin_per_unit")
    eng = _Engine([_hit("unit price is 9. unit variable cost is 5.", "f", {})])
    refused = _kernel_answer(eng, "c", "x", row, resolved=strict)
    assert refused.unknown is True
    assert refused.kernel_value is None
    assert f"guard failed: {guard}" in refused.text
    # Same corpus with no guard computes: the refusal came from the guard.
    assert _kernel_answer(eng, "c", "x", row, resolved=base).kernel_value == pytest.approx(4.0)


def _guard_overlay(guard: str):
    base = load_base_definitions()
    cmpu = next(d for d in base["definitions"] if d["id"] == "contribution_margin_per_unit")
    override = {
        **{k: cmpu[k] for k in ("id", "name", "category", "expression", "inputs", "output")},
        "definition_version": 1,
        "guards": [guard],
        "overrides": cmpu["key"],
        "reason": "Selling below variable cost is refused for this overlay.",
        "provenance": {"kind": "internal_protocol", "reference": "overlay guard test"},
    }
    return {"set_id": "overlay_guard_test", "definitions": [override]}


@pytest.mark.parametrize(
    "price, cost, overlay_refuses, value",
    [("5", "9", True, None), ("9", "5", False, 4.0)],
)
def test_answer_query_applies_an_overlay_guard(monkeypatch, price, cost, overlay_refuses, value):
    from app.cerebrum_product_kernel import answer as answer_mod

    monkeypatch.setattr(answer_mod.retrieval_index, "tenant_collection", lambda identity: "c-" + identity)
    guard = "unit_price > unit_variable_cost"
    eng = _Engine([_hit(f"unit price is {price}. unit variable cost is {cost}.", "f.md", {"line": 2})])
    q = "what is the contribution margin per unit"

    with_overlay = answer_mod.answer_query("t1", q, "ignored", engine=eng, overlays=[_guard_overlay(guard)])
    without = answer_mod.answer_query("t1", q, "ignored", engine=eng)
    assert with_overlay.kind == "kernel" and without.kind == "kernel"
    assert with_overlay.unknown is overlay_refuses
    assert without.unknown is False
    if overlay_refuses:
        assert f"guard failed: {guard}" in with_overlay.text
        assert without.kernel_value == pytest.approx(float(price) - float(cost))
    else:
        assert with_overlay.kernel_value == pytest.approx(value)
