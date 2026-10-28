"""F-1b: the parsers that turn a stored object into located chunks.

The binding rules asserted here, against the real kernel helpers (no mocks of
the parser itself, two genuine inputs):

* TWO-INPUT: a 2-page PDF and a 3-row CSV produce two DIFFERENT chunk sets; a
  PDF chunk carries ``location.page`` and a CSV chunk carries
  ``location.sheet`` + ``location.row``, distinct per file, and each chunk's
  text is derived from the file's own content.
* Per-type: PDF pages carry the right page numbers; XLSX chunks are tagged
  sheet + row; TXT / MD / DOCX chunks carry ``file`` + a structured location.
  Every parse moves the job ``queued`` -> ``parsed`` (asserted via get_job).
* A corrupt file of an accepted type lands the job at ``failed`` with
  ``last_error`` and the route does not 500.
* CROSS-TENANT: parsing tenant A's object neither reads nor writes under
  tenant B's digest.
"""

from __future__ import annotations

import io

import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient

from app.cerebrum_product_kernel.ingestion import object_store, parse_job
from app.cerebrum_product_kernel.ingestion.parse_job import (
    ParseStatus,
    create_job,
    get_job,
    read_chunks,
)
from app.cerebrum_product_kernel.ingestion.parsers import parse_object
from app.cerebrum_product_kernel.ingestion.router import router as ingestion_router
from app.factory.build.tenant_bind import bind_tenant_store


@pytest.fixture(autouse=True)
def _tenant_root(tmp_path, monkeypatch):
    monkeypatch.setenv("STORAGE_PATH", str(tmp_path / "storage"))
    monkeypatch.setenv("FACTORY_OUTPUTS_ROOT", str(tmp_path / "factory_outputs"))
    yield


# ── tiny real-document builders (valid enough for the real libraries) ────────


def _make_pdf(pages_text):
    """A minimal multi-page PDF whose pages carry extractable text (pypdf)."""
    n_pages = len(pages_text)
    parts = {}
    next_id = 3
    kids_ids = []
    page_to_content = {}
    content_text = {}
    for text in pages_text:
        page_id = next_id
        next_id += 1
        content_id = next_id
        next_id += 1
        kids_ids.append(page_id)
        page_to_content[page_id] = content_id
        content_text[content_id] = text

    parts[1] = b"<< /Type /Catalog /Pages 2 0 R >>"
    kids = " ".join(f"{k} 0 R" for k in kids_ids).encode()
    parts[2] = (
        b"<< /Type /Pages /Kids [" + kids + b"] /Count "
        + str(n_pages).encode() + b" >>"
    )
    for page_id, content_id in page_to_content.items():
        parts[page_id] = (
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            b"/Resources << /Font << /F1 << /Type /Font /Subtype /Type1 "
            b"/BaseFont /Helvetica >> >> >> /Contents "
            + str(content_id).encode() + b" 0 R >>"
        )
    for content_id, text in content_text.items():
        stream = b"BT /F1 24 Tf 72 700 Td (" + text.encode() + b") Tj ET"
        parts[content_id] = (
            b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n"
            + stream + b"\nendstream"
        )

    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets = {}
    for i in range(1, next_id):
        offsets[i] = out.tell()
        out.write(str(i).encode() + b" 0 obj\n" + parts[i] + b"\nendobj\n")
    xref_pos = out.tell()
    out.write(b"xref\n0 " + str(next_id).encode() + b"\n")
    out.write(b"0000000000 65535 f \n")
    for i in range(1, next_id):
        out.write(f"{offsets[i]:010d} 00000 n \n".encode())
    out.write(b"trailer\n<< /Size " + str(next_id).encode() + b" /Root 1 0 R >>\n")
    out.write(b"startxref\n" + str(xref_pos).encode() + b"\n%%EOF\n")
    return out.getvalue()


def _make_xlsx():
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Inventory"
    ws.append(["name", "qty"])
    ws.append(["widget", 5])
    ws.append(["gadget", 9])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _make_docx():
    import docx

    document = docx.Document()
    document.add_heading("Policy Section", level=1)
    document.add_paragraph("The first clause of the policy body.")
    document.add_paragraph("The second clause of the policy body.")
    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()


def _upload_and_parse(identity, source_name, data):
    stored = object_store.save_object(identity, source_name, data)
    job = create_job(
        identity, object_id=stored.object_id, source_name=stored.source_name
    )
    assert job.status is ParseStatus.QUEUED
    final = parse_object(identity, stored.object_id, job_id=job.job_id)
    return stored, final


# ── TWO-INPUT ────────────────────────────────────────────────────────────────


def test_two_inputs_two_located_chunk_sets():
    pdf_bytes = _make_pdf(["AlphaPageOne", "BravoPageTwo"])
    csv_bytes = b"name,qty\nwidget,5\ngadget,9\n"

    _, pdf_job = _upload_and_parse("tenant-a", "spec.pdf", pdf_bytes)
    _, csv_job = _upload_and_parse("tenant-a", "stock.csv", csv_bytes)

    assert pdf_job.status is ParseStatus.PARSED
    assert csv_job.status is ParseStatus.PARSED

    pdf_chunks = read_chunks("tenant-a", pdf_job.job_id)
    csv_chunks = read_chunks("tenant-a", csv_job.job_id)
    assert pdf_chunks and csv_chunks

    # Two DIFFERENT chunk sets.
    assert {c["id"] for c in pdf_chunks}.isdisjoint({c["id"] for c in csv_chunks})

    # A PDF chunk carries location.page; text is derived from the file content.
    page_two = [c for c in pdf_chunks if c["location"].get("page") == 2]
    assert page_two and "BravoPageTwo" in page_two[0]["text"]
    assert page_two[0]["file"] == "spec.pdf"
    assert "page" in page_two[0]["location"] and "row" not in page_two[0]["location"]

    # A CSV chunk carries location.sheet + location.row; text from the file.
    row_two = [c for c in csv_chunks if c["location"].get("row") == 2]
    assert row_two and "widget" in row_two[0]["text"]
    assert row_two[0]["location"]["sheet"] == "stock.csv"
    assert row_two[0]["file"] == "stock.csv"
    assert "page" not in row_two[0]["location"]


# ── per-type ──────────────────────────────────────────────────────────────────


def test_pdf_pages_carry_correct_page_numbers():
    pdf_bytes = _make_pdf(["FirstPageText", "SecondPageText", "ThirdPageText"])
    _, job = _upload_and_parse("tenant-a", "manual.pdf", pdf_bytes)
    assert job.status is ParseStatus.PARSED
    assert job.pages == 3

    chunks = read_chunks("tenant-a", job.job_id)
    by_page = {c["location"]["page"]: c for c in chunks}
    assert set(by_page) == {1, 2, 3}
    assert "FirstPageText" in by_page[1]["text"]
    assert "ThirdPageText" in by_page[3]["text"]
    assert all(c["file"] == "manual.pdf" for c in chunks)


def test_xlsx_chunks_tagged_sheet_and_row():
    _, job = _upload_and_parse("tenant-a", "book.xlsx", _make_xlsx())
    assert job.status is ParseStatus.PARSED
    assert job.sheets == 1
    assert job.rows == 3

    chunks = read_chunks("tenant-a", job.job_id)
    assert chunks
    for chunk in chunks:
        assert chunk["location"]["sheet"] == "Inventory"
        assert isinstance(chunk["location"]["row"], int)
        assert chunk["file"] == "book.xlsx"
    widget = [c for c in chunks if "widget" in c["text"]]
    assert widget and widget[0]["location"]["row"] == 2


def test_docx_chunks_carry_file_and_location():
    _, job = _upload_and_parse("tenant-a", "policy.docx", _make_docx())
    assert job.status is ParseStatus.PARSED

    chunks = read_chunks("tenant-a", job.job_id)
    assert chunks
    body = [c for c in chunks if "first clause" in c["text"]]
    assert body
    assert body[0]["file"] == "policy.docx"
    assert "paragraph" in body[0]["location"]
    # The nearest heading is threaded onto the body paragraph's location.
    assert body[0]["location"]["heading"] == "Policy Section"


@pytest.mark.parametrize("name", ["notes.txt", "readme.md"])
def test_text_and_md_chunks_carry_file_and_line(name):
    data = b"First paragraph line.\n\nSecond paragraph\nwith two lines.\n"
    _, job = _upload_and_parse("tenant-a", name, data)
    assert job.status is ParseStatus.PARSED

    chunks = read_chunks("tenant-a", job.job_id)
    assert chunks
    first = [c for c in chunks if "First paragraph" in c["text"]]
    second = [c for c in chunks if "Second paragraph" in c["text"]]
    assert first and second
    assert first[0]["file"] == name
    assert first[0]["location"]["line"] == 1
    assert second[0]["location"]["line"] == 3


def test_job_transitions_queued_to_parsed_observed_via_get_job():
    stored = object_store.save_object("tenant-a", "doc.md", b"Hello body text.\n")
    job = create_job(
        "tenant-a", object_id=stored.object_id, source_name=stored.source_name
    )
    assert get_job("tenant-a", job.job_id).status is ParseStatus.QUEUED
    parse_object("tenant-a", stored.object_id, job_id=job.job_id)
    assert get_job("tenant-a", job.job_id).status is ParseStatus.PARSED


# ── failed path ────────────────────────────────────────────────────────────────


def test_corrupt_accepted_type_fails_without_raising():
    # Looks like a .pdf (accepted type) but the bytes are not a real PDF.
    stored = object_store.save_object("tenant-a", "broken.pdf", b"%PDF-not-really")
    job = create_job(
        "tenant-a", object_id=stored.object_id, source_name=stored.source_name
    )
    final = parse_object("tenant-a", stored.object_id, job_id=job.job_id)
    assert final.status is ParseStatus.FAILED
    assert final.last_error
    # No chunks were persisted for the failed job.
    assert read_chunks("tenant-a", job.job_id) == []


def test_upload_route_does_not_500_on_corrupt_autoparsed_file(monkeypatch):
    """With auto-parse on, a corrupt accepted-type upload still answers 200 and
    the job is observably ``failed`` (never a 500)."""
    import sys
    import types

    monkeypatch.setenv("INGESTION_AUTO_PARSE", "1")
    mod = types.ModuleType("app.tenancy")

    class TenantRefused(Exception):
        pass

    class _Tenant:
        def __init__(self, tid):
            self.tenant_id = tid
            self.roles = []

    def resolve_tenant(headers):
        auth = (headers or {}).get("authorization", "")
        token = auth[7:].strip() if auth.startswith("Bearer ") else ""
        if token != "good-token":
            raise TenantRefused(token)
        return _Tenant("tenant-a")

    mod.TenantRefused = TenantRefused
    mod.resolve_tenant = resolve_tenant
    monkeypatch.setitem(sys.modules, "app.tenancy", mod)

    app = FastAPI()
    app.include_router(ingestion_router)
    client = TestClient(app)
    headers = {"Authorization": "Bearer good-token"}

    resp = client.post(
        "/v1/product/ingestion/upload",
        headers=headers,
        files={"files": ("broken.pdf", b"%PDF-not-really", "application/pdf")},
    )
    assert resp.status_code == 200, resp.text
    assert resp.json()["status"] == "failed"

    status = client.get(
        f"/v1/product/ingestion/upload/{resp.json()['job_id']}", headers=headers
    )
    assert status.status_code == 200
    assert status.json()["status"] == "failed"
    assert status.json()["last_error"]


def test_autoparse_reaches_parsed_through_the_route(monkeypatch):
    """A real multi-page PDF uploaded with auto-parse on reaches ``parsed``."""
    import sys
    import types

    monkeypatch.setenv("INGESTION_AUTO_PARSE", "1")
    mod = types.ModuleType("app.tenancy")

    class TenantRefused(Exception):
        pass

    class _Tenant:
        def __init__(self, tid):
            self.tenant_id = tid
            self.roles = []

    def resolve_tenant(headers):
        auth = (headers or {}).get("authorization", "")
        token = auth[7:].strip() if auth.startswith("Bearer ") else ""
        if token != "good-token":
            raise TenantRefused(token)
        return _Tenant("tenant-a")

    mod.TenantRefused = TenantRefused
    mod.resolve_tenant = resolve_tenant
    monkeypatch.setitem(sys.modules, "app.tenancy", mod)

    app = FastAPI()
    app.include_router(ingestion_router)
    client = TestClient(app)
    headers = {"Authorization": "Bearer good-token"}

    resp = client.post(
        "/v1/product/ingestion/upload",
        headers=headers,
        files={"files": ("real.pdf", _make_pdf(["One", "Two"]), "application/pdf")},
    )
    assert resp.status_code == 200, resp.text
    assert resp.json()["status"] == "parsed"
    status = client.get(
        f"/v1/product/ingestion/upload/{resp.json()['job_id']}", headers=headers
    )
    assert status.json()["status"] == "parsed"
    assert status.json()["pages"] == 2


# ── cross-tenant isolation ────────────────────────────────────────────────────


def test_parsing_tenant_a_never_touches_tenant_b():
    _, job = _upload_and_parse("tenant-a", "secret.csv", b"h\nsecret-row\n")
    assert job.status is ParseStatus.PARSED

    # The chunks landed under A's own digest dir, never under B's.
    a_chunks_path = parse_job.chunks_path("tenant-a", job.job_id)
    assert a_chunks_path.is_file()
    a_key = bind_tenant_store("tenant-a").tenant_key
    b_key = bind_tenant_store("tenant-b").tenant_key
    assert a_key in str(a_chunks_path)
    assert b_key not in str(a_chunks_path)

    # B has no view of A's job or chunks, and cannot resolve a parse for it.
    assert get_job("tenant-b", job.job_id) is None
    assert read_chunks("tenant-b", job.job_id) == []
    with pytest.raises(ValueError):
        parse_object("tenant-b", job.object_id)
    # A's object is not even readable under B's digest.
    with pytest.raises(object_store.ObjectStoreError):
        object_store.read_object("tenant-b", job.object_id)
